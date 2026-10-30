from __future__ import annotations

import os
import tempfile
from pathlib import Path
from typing import Iterable

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = ROOT / "docs" / "competition_submission"
ASSET_DIR = Path(tempfile.gettempdir()) / "ai_competition_document_assets"

TEAM_NAME = "审查队"
PROJECT_NAME = "面向工程设计方案的行业标准大模型智能体审查与风险预警系统"
MEMBERS = "涂逸焯、张海鹏、何壮壮（均为硕士研究生）"
VERSION = "V1.0"
DATE = "2026.07.22"
GROUP = "待报名系统确认"

SHORT_TITLE = "行业标准大模型智能体审查与风险预警系统"
BODY_FONT = "宋体"
HEADING_FONT = "黑体"
LATIN_FONT = "Times New Roman"
USABLE_DXA = 9060

INTRO_TEXT = (
    "本项目面向工程设计成果中标准资料分散、人工审查耗时、空间风险难定位等问题，构建行业标准大模型智能体"
    "审查与风险预警系统。系统以GIS设计成果为直接审查对象，通过GeoPandas与Pydantic完成多源数据适配和"
    "工程对象语义标准化，利用Shapely计算空间关系，并由通用GIS数据质量、项目业务及国内外标准规范三类规则库"
    "形成稳定、可追溯的风险结论。bge-m3召回相关条款，LangGraph编排本地32B模型、QGIS MCP和证据校验，"
    "生成依据明确、可执行的整改建议，实现风险定位、可视化分析及多格式审查报告导出，为不同地区和工程类型"
    "提供可扩展的智能审查能力。"
)


def set_run_font(run, name=BODY_FONT, size=11, bold=None, color=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), LATIN_FONT if name == BODY_FONT else name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), LATIN_FONT if name == BODY_FONT else name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent=120):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            cell.width = Twips(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, headers: list[str], rows: Iterable[Iterable[str]], widths: list[int], font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.rows[0].cells
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(header), HEADING_FONT, font_size, bold=True)
        set_cell_shading(cell, "EDEDED")
    set_repeat_table_header(table.rows[0])
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if len(str(value)) < 18 else WD_ALIGN_PARAGRAPH.LEFT
            set_run_font(p.add_run(str(value)), BODY_FONT, font_size)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def configure_styles(doc: Document):
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), LATIN_FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), LATIN_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5

    for name, size, before, after in (
        ("Heading 1", 16, 14, 8),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 12, 9, 5),
    ):
        style = doc.styles[name]
        style.font.name = HEADING_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:ascii"), HEADING_FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), HEADING_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for list_style in ("List Bullet", "List Number"):
        style = doc.styles[list_style]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Cm(0.74)
        style.paragraph_format.first_line_indent = Cm(-0.37)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.35


def configure_section(section):
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)


def add_bottom_border(paragraph, color="333333", size="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)


def set_header_footer(section, footer_text=None, body_page=False):
    section.header.is_linked_to_previous = False
    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = ""
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_p.paragraph_format.space_after = Pt(2)
    set_run_font(header_p.add_run(SHORT_TITLE), BODY_FONT, 9, color="555555")
    add_bottom_border(header_p, "444444", "6")

    section.footer.is_linked_to_previous = False
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = ""
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if body_page:
        run = footer_p.add_run()
        begin = OxmlElement("w:fldChar")
        begin.set(qn("w:fldCharType"), "begin")
        instr = OxmlElement("w:instrText")
        instr.set(qn("xml:space"), "preserve")
        instr.text = " PAGE "
        separate = OxmlElement("w:fldChar")
        separate.set(qn("w:fldCharType"), "separate")
        text = OxmlElement("w:t")
        text.text = "1"
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        for node in (begin, instr, separate, text, end):
            run._r.append(node)
        set_run_font(run, LATIN_FONT, 9)
        sect_pr = section._sectPr
        pg_num = sect_pr.find(qn("w:pgNumType"))
        if pg_num is None:
            pg_num = OxmlElement("w:pgNumType")
            sect_pr.append(pg_num)
        pg_num.set(qn("w:start"), "1")
        pg_num.set(qn("w:fmt"), "decimal")
    elif footer_text:
        set_run_font(footer_p.add_run(footer_text), LATIN_FONT, 9)


def add_section(doc, footer_text=None, body_page=False):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(section)
    set_header_footer(section, footer_text=footer_text, body_page=body_page)
    return section


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.widow_control = True
    if bold_prefix and text.startswith(bold_prefix):
        set_run_font(p.add_run(bold_prefix), BODY_FONT, 11, bold=True)
        set_run_font(p.add_run(text[len(bold_prefix) :]), BODY_FONT, 11)
    else:
        set_run_font(p.add_run(text), BODY_FONT, 11)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_run_font(p.add_run(text), BODY_FONT, 11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    set_run_font(p.add_run(text), BODY_FONT, 11)
    return p


def add_heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    set_run_font(p.add_run(text), HEADING_FONT, {1: 16, 2: 14, 3: 12}[level], bold=True)
    return p


def add_toc_field(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-2" \\h \\z \\u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目录将在打开文档时自动更新"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)
    return p


def wrap_text(draw, text, font, max_width):
    lines = []
    current = ""
    for char in text:
        candidate = current + char
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = char
    if current:
        lines.append(current)
    return lines


def draw_box(draw, xy, title, subtitle, font_title, font_sub, fill, outline):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=16, fill=fill, outline=outline, width=3)
    title_lines = wrap_text(draw, title, font_title, x2 - x1 - 30)
    subtitle_lines = wrap_text(draw, subtitle, font_sub, x2 - x1 - 30)
    total = len(title_lines) * 36 + len(subtitle_lines) * 28 + 8
    y = y1 + (y2 - y1 - total) / 2
    for line in title_lines:
        bbox = draw.textbbox((0, 0), line, font=font_title)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=font_title, fill="#14212b")
        y += 36
    y += 8
    for line in subtitle_lines:
        bbox = draw.textbbox((0, 0), line, font=font_sub)
        draw.text(((x1 + x2 - (bbox[2] - bbox[0])) / 2, y), line, font=font_sub, fill="#405666")
        y += 28


def create_architecture_diagram(path: Path):
    image = Image.new("RGB", (1500, 1050), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 28)
    sub_font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 20)
    arrow_font = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 34)
    boxes = [
        ((400, 35, 1100, 155), "GIS工程设计成果", "14个GPKG / QGZ / GeoJSON / Shapefile", "#EAF4F1", "#147D6F"),
        ((400, 205, 1100, 325), "数据适配与语义标准化", "GeoPandas / Pyogrio读取，Pydantic统一对象模型", "#EDF3FA", "#2F6EA4"),
        ((400, 375, 1100, 495), "GIS空间特征计算", "Shapely：距离、相交、包含、穿越、几何有效性", "#F6F1E8", "#A96F15"),
        ((40, 555, 470, 690), "三类规则库", "通用GIS质量 / 项目业务 / 德国与中国技术规范", "#F9EEEE", "#B94A48"),
        ((535, 555, 965, 690), "确定性规则引擎", "输出正式风险、计算证据和风险位置", "#F9EEEE", "#B94A48"),
        ((1030, 555, 1460, 690), "LangGraph智能体", "bge-m3召回 / 32B解释 / MCP工具 / 证据校验", "#F1EEFA", "#7555A7"),
        ((300, 790, 1200, 940), "风险定位、预警与审查报告", "WebGIS/QGIS展示，PDF/Word/Excel/GeoJSON导出，可追溯整改建议", "#EAF4F1", "#147D6F"),
    ]
    for box in boxes:
        draw_box(draw, *box[:3], title_font, sub_font, box[3], box[4])
    for x, y, symbol in ((750, 170, "↓"), (750, 340, "↓"), (750, 510, "↓"), (500, 620, "→"), (995, 620, "→"), (750, 720, "↓")):
        draw.text((x - 16, y), symbol, font=arrow_font, fill="#526571")
    image.save(path)


def create_langgraph_diagram(path: Path):
    image = Image.new("RGB", (1500, 900), "white")
    draw = ImageDraw.Draw(image)
    title_font = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 25)
    sub_font = ImageFont.truetype(r"C:\Windows\Fonts\msyh.ttc", 18)
    arrow_font = ImageFont.truetype(r"C:\Windows\Fonts\msyhbd.ttc", 30)
    nodes = [
        ("运行状态检查", "runtime_status"),
        ("数据画像", "data_profile"),
        ("语义检索", "semantic_retrieval"),
        ("GIS工具规划", "gis_tool_planning"),
        ("MCP/GIS工具执行", "mcp_gis_tools"),
        ("证据校验", "evidence_validation"),
    ]
    x_positions = [25, 270, 515, 760, 1005, 1250]
    for x, (title, subtitle) in zip(x_positions, nodes):
        draw_box(draw, (x, 125, x + 220, 275), title, subtitle, title_font, sub_font, "#F2F7F8", "#278679")
    for x in (245, 490, 735, 980, 1225):
        draw.text((x - 11, 178), "→", font=arrow_font, fill="#526571")
    draw.text((720, 295), "↓", font=arrow_font, fill="#526571")
    draw_box(draw, (430, 350, 700, 500), "本地32B模型审查", "llm_review", title_font, sub_font, "#EEF3FA", "#3D6F9E")
    draw_box(draw, (800, 350, 1070, 500), "规则结果回退", "fallback_review", title_font, sub_font, "#FAF2E8", "#A87022")
    draw.text((705, 405), "或", font=title_font, fill="#526571")
    draw.text((550, 525), "↓", font=arrow_font, fill="#526571")
    draw.text((920, 525), "↓", font=arrow_font, fill="#526571")
    draw_box(draw, (480, 595, 1020, 745), "统一智能审查结论", "模型失败不影响正式规则风险输出", title_font, sub_font, "#F7EEEE", "#B94A48")
    image.save(path)


def create_project_document(output_path: Path):
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_section(section)
    set_header_footer(section, footer_text="I")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(72)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("第八届中国研究生人工智能创新大赛"), BODY_FONT, 22, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(120)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(PROJECT_NAME), HEADING_FONT, 20, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(24)
    set_run_font(p.add_run("项目文档"), HEADING_FONT, 18, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run(VERSION), LATIN_FONT, 13)
    for text in (DATE, TEAM_NAME, f"团队成员：{MEMBERS}", f"参赛组别：{GROUP}"):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4 if text != DATE else 70)
        p.paragraph_format.space_after = Pt(0)
        set_run_font(p.add_run(text), BODY_FONT, 12, bold=text in (TEAM_NAME,))

    add_section(doc, footer_text="II")
    p = doc.add_paragraph()
    set_run_font(p.add_run("目录"), HEADING_FONT, 22, bold=True)
    p.paragraph_format.space_after = Pt(20)
    add_toc_field(doc)

    add_section(doc, footer_text="III")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("记录更改历史"), HEADING_FONT, 18, bold=True)
    rows = [
        ("1", "形成初赛项目文档", VERSION, "审查队", DATE, "基于当前可运行原型"),
        ("", "", "", "", "", ""),
        ("", "", "", "", "", ""),
        ("", "", "", "", "", ""),
        ("", "", "", "", "", ""),
    ]
    add_table(doc, ["序号", "更改原因", "版本", "作者", "更改日期", "备注"], rows, [650, 2100, 900, 1300, 1500, 2610], 9.5)

    add_section(doc, body_page=True)
    add_heading(doc, "1 项目概况", 1)
    add_heading(doc, "1.1 背景和基础", 2)
    add_body(
        doc,
        "通信管网、地下管线和基建工程设计通常同时包含GIS图层、属性表、技术说明与多来源行业标准。"
        "现行审查工作高度依赖工程师逐图层核对、人工查找条款和经验判断，存在标准分散、跨语言理解成本高、"
        "空间关系难批量计算、审查结果难复现以及整改依据难追溯等问题。通用大模型虽能生成文本建议，"
        "但如果直接承担合规判定，容易受到随机性和幻觉影响，难以满足工程审查的确定性要求。",
    )
    add_body(
        doc,
        "项目由涂逸焯、张海鹏、何壮壮三名硕士研究生组成。团队围绕“数据适配与语义标准化—GIS空间计算—"
        "规则适用性判断—大模型解释—风险定位与报告输出”开展研发，已经形成可运行的Web原型，并完成德国"
        "Dreieich光纤工程样例的端到端验证。",
    )
    add_body(
        doc,
        "当前基础包括：Vue 3/TypeScript前端、FastAPI后端、Pydantic统一对象模型、GeoPandas/Pyogrio数据读取、"
        "Shapely空间计算、三类后端规则库、bge-m3标准条款检索、qwen2.5:32b本地模型、LangGraph智能体工作流、"
        "Docker化QGIS Server和QGIS MCP工具服务。正式风险由确定性规则引擎输出，大模型只承担工具规划、解释和"
        "整改建议补充，从架构上保证同一输入的正式风险结果稳定。",
    )

    add_heading(doc, "1.2 场景和价值", 2)
    add_body(
        doc,
        "系统适用于通信管道、光纤接入、城市地下管线、道路开挖及其他以GIS设计成果为主要交付物的工程审查场景。"
        "用户上传GPKG、GeoJSON或Shapefile后，系统自动读取图层、标准化工程对象、计算空间特征、执行规则并在地图"
        "定位风险，审查人员可进一步查看原始对象、计算证据、适用条款和整改路径。",
    )
    add_table(
        doc,
        ["方案", "优势", "主要局限", "本项目改进"],
        [
            ("人工/QGIS审查", "专业判断强，图形表达准确", "耗时、依赖经验、难批量复现", "自动批量计算并保留人工复核入口"),
            ("纯规则系统", "结果稳定、可解释", "规则适配成本高，文本说明能力弱", "三类规则库结合标准检索与智能体解释"),
            ("通用大模型/RAG", "自然语言理解和生成能力强", "可能幻觉，空间计算与风险计数不稳定", "模型不决定正式风险，工具证据先校验后使用"),
            ("本项目", "规则确定性、GIS计算、标准RAG与智能体协同", "当前仍需扩充标注集与行业验证", "形成可追溯、可定位、可复现的闭环"),
        ],
        [1350, 2250, 2500, 2960],
        9,
    )
    add_body(
        doc,
        "潜在价值体现在四方面：一是降低海量要素逐项核查的人力成本；二是用固定规则和证据链提升审查一致性；"
        "三是通过风险位置聚合和QGIS定位缩短整改闭环；四是采用本地模型和后端只读标准库，兼顾工程数据安全与"
        "审查依据治理。",
    )

    add_heading(doc, "1.3 所需支持", 2)
    for item in (
        "算力支持：稳定运行32B参数级模型的GPU服务器，以及bge-m3向量检索和模型推理的并发资源。",
        "行业支持：通信管道、道路开挖、地下管线等领域专家对规则适用范围、阈值和整改建议进行复核。",
        "数据支持：更多经过脱敏和授权的真实GIS设计成果，以及可用于准确率评估的人工标注审查集。",
        "工程支持：容器部署、性能压测、模型量化与缓存优化培训，以及比赛期间的云端演示环境。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "2 项目规划", 1)
    add_heading(doc, "2.1 整体目标", 2)
    add_body(
        doc,
        "参赛期间目标是形成可公开演示、可重复运行、可量化评估的行业智能审查原型。系统应完成GIS成果上传、"
        "对象标准化、空间特征计算、三类规则审查、标准语义召回、智能体工具调用、风险地图定位、整改建议和多格式"
        "报告导出，并在至少一个德国工程样例和一个中国标准适配样例上进行验证。",
    )
    add_table(
        doc,
        ["目标维度", "当前状态", "竞赛期目标"],
        [
            ("功能闭环", "主流程已跑通", "形成稳定演示脚本和异常回退流程"),
            ("确定性", "正式风险由规则引擎固定输出", "同输入重复审查的正式风险数量与等级一致率100%"),
            ("证据追溯", "支持对象、条款、计算证据和来源展示", "抽样风险证据可追溯率不低于95%"),
            ("规则效果", "24条条款、42条可执行规则", "建立人工标注测试集，目标规则判定准确率不低于90%"),
            ("性能", "万级对象可完成审查，模型阶段仍需优化", "通过缓存、裁剪和并发使演示任务端到端控制在5分钟级"),
            ("跨标准适配", "德国9份规则资料+中国5份标准资料", "形成标准到对象字段/空间谓词的半自动映射流程"),
        ],
        [1450, 3400, 4210],
        9,
    )

    add_heading(doc, "2.2 技术创新点", 2)
    innovations = [
        "面向GIS成果的语义标准化层：不是把PDF或设计说明直接交给大模型，而是先把多图层、多语言字段映射为统一工程对象，为后续规则和空间计算提供稳定输入。",
        "三类规则库协同：将通用GIS数据质量、项目数据与业务、德国/中国标准与技术规范分层管理，兼顾可迁移规则、项目特有规则和法规依据。",
        "确定性风险与生成式解释解耦：风险数量、等级和阈值证据由固定规则引擎决定；32B模型负责解释、工具规划和整改表达，避免同一方案多次审查结果漂移。",
        "LangGraph证据门控工作流：通过数据画像、语义检索、GIS工具规划、MCP工具执行和证据校验节点，只有通过校验的GIS事实才进入大模型结论。",
        "QGIS原生渲染与MCP整改闭环：QGIS Server复用项目样式呈现工程地图，QGIS MCP验证工程和图层上下文，将风险对象定位转化为可操作的整改路径。",
        "风险位置聚合：同一对象命中多条规则时按最高风险等级归并为一个位置卡片，同时保留全部规则构成，降低地图标记和清单重复。",
    ]
    for item in innovations:
        add_bullet(doc, item)

    add_heading(doc, "3 实施方案", 1)
    add_heading(doc, "3.1 技术可行性分析", 2)
    add_heading(doc, "3.1.1 数据与行业知识来源", 3)
    add_body(
        doc,
        "工程数据采用赛题/项目方提供的德国Dreieich光纤工程样例，包含14个GPKG、1个QGIS工程及配套照片资料。"
        "规则知识来源包括德国行业标准文件夹中的10份PDF（其中00号文档为规则适配说明，其余9份为技术规则资料）"
        "以及中国行业标准文件夹中的5份PDF。当前后端已形成24条结构化条款和42条可执行规则。所有标准在正式库中"
        "记录来源文件、条款编号、摘要、适用对象、风险等级、规则参数和整改建议，网页端默认只读。",
    )
    add_heading(doc, "3.1.2 算法与软件来源", 3)
    add_body(
        doc,
        "系统使用成熟开源组件完成工程能力：GeoPandas/Pyogrio读取GPKG、GeoJSON和Shapefile；Pydantic定义统一"
        "工程对象、规则与审查结果；Shapely计算距离、相交、包含、穿越和几何有效性；FastAPI提供接口；Vue 3、"
        "OpenLayers和ECharts构建交互式WebGIS；LangGraph编排状态、节点与条件边；Ollama承载本地qwen2.5:32b和"
        "bge-m3；QGIS Server复用QGIS项目渲染，MCP服务向智能体提供只读工程能力校验。相关来源列于第4章。",
    )
    add_heading(doc, "3.1.3 算力与部署条件", 3)
    add_body(
        doc,
        "当前原型运行于Windows和ziketi3虚拟环境，FastAPI监听本地8010端口；模型服务部署于团队可访问的计算"
        "服务器并通过本地接口调用；QGIS Server与QGIS MCP以Docker容器运行。即使32B模型暂时不可用，LangGraph"
        "也会进入fallback_review，规则引擎仍可输出正式风险，因此系统具备可降级运行能力。后续将按实际设备补充"
        "GPU型号、显存和功耗信息，并开展推理时延与并发测试。",
    )

    add_heading(doc, "3.2 技术细节", 2)
    create_architecture_diagram(ASSET_DIR / "architecture.png")
    doc.add_picture(str(ASSET_DIR / "architecture.png"), width=Cm(15.5))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("图1 系统总体技术架构"), BODY_FONT, 10.5)

    add_heading(doc, "3.2.1 数据适配与语义标准化", 3)
    add_body(
        doc,
        "后端遍历GIS文件及其图层，读取几何、字段、坐标参考系和来源信息，并映射为DesignObject。对象至少包含"
        "id、type、name、geometry、properties和source。德文字段和工程图层名称通过项目输入画像与映射规则转换为"
        "统一对象类型，例如telecom_duct_route、civil_pit_work和fiber_distribution_cabinet，使标准规则不依赖"
        "某一文件的原始命名。",
    )
    add_heading(doc, "3.2.2 GIS空间特征与规则执行", 3)
    add_body(
        doc,
        "Shapely将GeoJSON几何转换为可计算对象，规则引擎执行required_geometry、geometry_validity、"
        "required_properties、min_property、min_distance、forbidden_intersection、topology_connectivity等规则。"
        "每条正式风险保存对象ID、几何、实际值、阈值、单位、标准条款和计算表达式。规则执行日志同时记录可计算性"
        "与命中数量，支持人工复核。",
    )
    add_heading(doc, "3.2.3 LangGraph智能体与证据门控", 3)
    create_langgraph_diagram(ASSET_DIR / "langgraph.png")
    doc.add_picture(str(ASSET_DIR / "langgraph.png"), width=Cm(15.5))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("图2 LangGraph智能审查工作流"), BODY_FONT, 10.5)
    add_body(
        doc,
        "工作流状态保存工程对象、正式风险、执行日志、数据画像、召回条款、工具计划、工具结果、运行状态和证据"
        "校验结果。runtime_status检查模型服务；data_profile统计对象与几何；semantic_retrieval使用bge-m3召回标准；"
        "gis_tool_planning由本地模型生成白名单内的只读工具计划；mcp_gis_tools执行QGIS MCP和本地GIS工具；"
        "evidence_validation校验对象编号和结果结构；最后根据模型可用性进入llm_review或fallback_review。",
    )
    add_heading(doc, "3.2.4 QGIS地图与整改闭环", 3)
    add_body(
        doc,
        "QGIS Server读取原始QGZ项目并通过WMS输出地图，OpenLayers在网页端叠加风险位置。用户点击审查结果后，"
        "地图定位对应对象；同一位置的多条规则按最高等级归类并在详情中展开。QGIS MCP验证工程可用性和对象图层"
        "上下文，整改建议给出“进入对应工程图层—筛选对象—修改字段或几何—保存并重新审查”的操作路径。",
    )
    add_heading(doc, "3.2.5 当前推理效果与指标", 3)
    add_body(
        doc,
        "以下为当前德国样例的一次完整运行记录，反映工程可运行性，不等同于经人工标注得到的准确率或召回率。"
        "正式精度指标将在竞赛期通过人工标注测试集计算。",
    )
    add_table(
        doc,
        ["指标", "当前结果", "说明"],
        [
            ("输入规模", "14个GPKG，12406个对象", "德国Dreieich工程样例"),
            ("规则知识", "24条结构化条款，42条可执行规则", "三类规则库合计"),
            ("规则命中", "1376条", "高风险883条，中风险493条"),
            ("风险位置", "740个", "同一对象多规则按最高等级聚合"),
            ("语义召回", "24条标准证据", "bge-m3本次运行召回结果"),
            ("智能体状态", "LangGraph运行完成", "QGIS MCP已连接，证据校验后进入模型节点"),
            ("结果稳定性", "同一样例的正式风险计数保持一致", "正式风险由确定性规则引擎产生"),
        ],
        [1800, 2550, 4710],
        9.5,
    )
    add_body(
        doc,
        "当前局限包括：测试样例仍以单一德国项目为主；部分标准条款需要行业专家将自然语言进一步转化为可计算"
        "约束；32B模型推理耗时受服务器负载影响；尚未建立覆盖全部规则类型的金标准测试集。系统已通过确定性规则"
        "与模型回退机制控制上述风险，后续重点开展标注评测、性能优化和跨项目验证。",
    )

    add_heading(doc, "3.3 计划和分工", 2)
    add_table(
        doc,
        ["阶段", "时间", "主要任务", "责任分工", "交付物"],
        [
            ("基线固化", "7月22日-7月31日", "冻结原型版本，补齐测试与材料来源说明", "涂逸焯统筹；三人交叉复核", "可重复演示版本"),
            ("规则与评测", "8月1日-8月10日", "扩充标注集，评估规则准确率和证据追溯", "张海鹏主责；何壮壮协助", "评测报告与测试集"),
            ("智能体优化", "8月11日-8月18日", "优化提示、检索、MCP工具与推理时延", "涂逸焯主责；张海鹏复核规则", "稳定智能体流程"),
            ("前端与展示", "8月19日-8月25日", "完善地图交互、报告和异常提示", "何壮壮主责；涂逸焯联调", "演示网站与报告"),
            ("材料定稿", "8月26日-8月31日", "核对文档、简介、代码清单和提交格式", "三人共同完成", "正式提交包"),
        ],
        [1150, 1500, 2850, 1900, 1660],
        8.7,
    )
    add_body(doc, "建议团队模块分工如下，最终可按实际贡献调整：")
    for item in (
        "涂逸焯：总体架构、LangGraph智能体、本地模型与QGIS MCP集成、后端联调。",
        "张海鹏：GIS数据适配、标准知识与规则库、空间规则验证、测试集设计。",
        "何壮壮：Vue/WebGIS交互、风险统计与报告展示、系统测试、参赛材料整合。",
        "共同任务：行业调研、专家复核、演示验证、文档与答辩准备。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "4 参考资料", 1)
    references = [
        "第八届中国研究生人工智能创新大赛初赛作品提交规范，2025-05-22。",
        "第八届中国研究生人工智能创新大赛项目文档模板。",
        "Richtlinie für das Aufgraben von Straßen, Wegen und Plätzen in der Stadt Dreieich, 2014。",
        "GIS-Nebenbestimmungen Version 5.1, 2023。",
        "Materialkonzept Version 5.0.2。",
        "Einheitliches Messkonzept, 2022。",
        "OEW光纤网络规划与文档指南 V1.1, 2025。",
        "德国电信电缆保护指令, 2026。",
        "DGUV Information 203-017, 2019。",
        "NRM-N-A001地下公用管线保护规范, 2023。",
        "GB 50373-2019 通信管道与通道工程设计标准。",
        "GB/T 50374-2018 通信管道工程施工及验收标准。",
        "DB4403/T 587-2025 深圳市地下管线数据规范。",
        "广东省建筑物配建通信基础设施实施指南（试行）。",
        "中兴家庭光纤施工技术白皮书。",
        "LangGraph官方文档，https://docs.langchain.com/oss/python/langgraph/overview。",
        "GeoPandas官方文档，https://geopandas.org/en/stable/docs/user_guide/io.html。",
        "Shapely官方文档，https://shapely.readthedocs.io/en/stable/manual.html。",
        "Pydantic官方文档，https://docs.pydantic.dev/latest/。",
        "FastAPI官方文档，https://fastapi.tiangolo.com/。",
        "Ollama Embeddings官方文档，https://docs.ollama.com/capabilities/embeddings。",
        "QGIS Server官方文档，https://docs.qgis.org/3.44/en/docs/server_manual/。",
    ]
    for ref in references:
        add_number(doc, ref)

    doc.core_properties.title = PROJECT_NAME
    doc.core_properties.subject = "第八届中国研究生人工智能创新大赛项目文档"
    doc.core_properties.author = TEAM_NAME
    doc.core_properties.keywords = "人工智能, GIS, 行业标准, LangGraph, QGIS, 风险预警"
    doc.save(output_path)


def create_intro_document(output_path: Path):
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_section(section)
    section.top_margin = Cm(2.6)
    section.bottom_margin = Cm(2.6)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(28)
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run("第八届中国研究生人工智能创新大赛"), HEADING_FONT, 18, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    set_run_font(p.add_run("参赛作品简介"), HEADING_FONT, 22, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(24)
    set_run_font(p.add_run(PROJECT_NAME), BODY_FONT, 14, bold=True)

    meta = add_table(
        doc,
        ["团队名称", "团队成员"],
        [(TEAM_NAME, MEMBERS)],
        [2100, 6960],
        10.5,
    )
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(10)
    set_run_font(p.add_run("作品简介"), HEADING_FONT, 14, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.65
    p.paragraph_format.space_after = Pt(18)
    set_run_font(p.add_run(INTRO_TEXT), BODY_FONT, 12)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    set_run_font(p.add_run("关键词："), BODY_FONT, 10.5, bold=True)
    set_run_font(p.add_run("工程设计审查；行业标准；GIS；大模型智能体；LangGraph；风险预警"), BODY_FONT, 10.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(28)
    set_run_font(p.add_run(f"{TEAM_NAME}  |  {DATE}"), BODY_FONT, 10, color="555555")
    doc.core_properties.title = f"{PROJECT_NAME}参赛作品简介"
    doc.core_properties.author = TEAM_NAME
    doc.save(output_path)


def create_auxiliary_document(output_path: Path):
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    configure_section(section)
    set_header_footer(section, body_page=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(35)
    set_run_font(p.add_run("初赛提交材料清单与技术可行性说明"), HEADING_FONT, 20, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    set_run_font(p.add_run(PROJECT_NAME), BODY_FONT, 13, bold=True)
    add_table(
        doc,
        ["项目", "内容"],
        [
            ("团队", f"{TEAM_NAME}；成员：{MEMBERS}"),
            ("版本", f"{VERSION}；{DATE}"),
            ("用途", "作为“其他辅助材料”的说明文件，最终可与源代码、环境文件和测试样例清单一起打包。"),
        ],
        [1500, 7560],
        10,
    )

    add_heading(doc, "1 初赛材料清单", 1)
    add_table(
        doc,
        ["材料", "附件要求", "当前状态", "正式提交前操作"],
        [
            ("参赛作品简介", "300字以内，PDF", "已生成Word稿", "复核参赛组别后转PDF并按规范命名"),
            ("项目文档", "基于附件2模板，PDF", "已生成Word稿", "更新目录、核对团队信息后转PDF"),
            ("项目视频", "不超过5分钟、200MB、MP4", "按团队要求暂不生成", "后续录制核心目标与预期效果"),
            ("其他辅助材料", "ZIP，不超过200MB", "本说明可纳入", "加入代码、README、环境清单和测试说明后压缩"),
        ],
        [1500, 2400, 2200, 2960],
        9.3,
    )
    add_body(doc, "提交截止时间：2026年9月1日23:59。学生参赛资格需由研究生培养单位审核通过。")
    add_body(doc, "正式文件建议命名：")
    for item in (
        f"{TEAM_NAME}_{PROJECT_NAME}_参赛作品简介.pdf",
        f"{TEAM_NAME}_{PROJECT_NAME}_项目文档.pdf",
        f"{TEAM_NAME}_{PROJECT_NAME}_项目视频.mp4",
        f"{TEAM_NAME}_{PROJECT_NAME}_其他.zip",
    ):
        add_bullet(doc, item)

    add_heading(doc, "2 技术可行性摘要", 1)
    add_body(
        doc,
        "系统已经完成从GIS成果上传到风险报告导出的端到端闭环。前端使用Vue 3、TypeScript、OpenLayers和"
        "ECharts；后端使用FastAPI、Pydantic、GeoPandas/Pyogrio和Shapely；智能体使用LangGraph编排本地"
        "qwen2.5:32b、bge-m3及QGIS MCP；QGIS Server负责复用原工程样式渲染。正式风险由确定性规则引擎输出，"
        "模型不可用时仍能完成基础审查。",
    )
    add_table(
        doc,
        ["可行性维度", "已有条件", "待完善事项"],
        [
            ("数据", "德国工程14个GPKG、12406个对象", "扩充脱敏项目和人工标注集"),
            ("知识", "德国9份规则资料、中国5份标准资料，24条结构化条款", "专家复核适用性与阈值"),
            ("规则", "42条可执行规则，覆盖属性、几何、距离、拓扑等", "增加跨图层组合规则和边界案例"),
            ("模型", "qwen2.5:32b与bge-m3本地接口可用", "优化时延、缓存和模型输出评测"),
            ("GIS", "GeoPandas/Shapely计算，QGIS Server/MCP已接入", "完善工具白名单和空间证据校验"),
            ("工程化", "FastAPI+Vue生产构建，支持多格式报告", "容器编排、日志监控和并发压测"),
        ],
        [1650, 4150, 3260],
        9.3,
    )

    add_heading(doc, "3 其他ZIP建议目录", 1)
    for item in (
        "01_项目源代码：review_system、frontend、scripts及必要配置示例；排除.env、模型权重、日志和临时项目数据。",
        "02_运行环境：requirements.txt、environment.yml、Docker/QGIS服务说明、启动与停止命令。",
        "03_规则与知识说明：三类规则库目录、标准来源清单、规则字段说明和只读治理策略。",
        "04_测试与指标：德国样例输入清单、测试集设计、规则命中统计、风险位置统计和已知局限。",
        "05_展示材料：系统截图、架构图、LangGraph流程图和示例审查报告。",
    ):
        add_number(doc, item)

    add_heading(doc, "4 提交前合规检查", 1)
    for item in (
        "确认参赛组别、培养单位、指导教师及成员排序与报名系统一致。",
        "将两份Word主文档转换为PDF，检查目录页码、字体、图片和表格无错位。",
        "核对简介正文不超过300字，项目名称和团队名称在全部文件中完全一致。",
        "在项目文档中保留数据、行业知识、算法和硬件来源，不写无法验证的准确率或性能结论。",
        "ZIP中不得包含.env、真实密钥、服务器账号、原始敏感工程数据或未经授权的标准全文。",
        "代码包应能按README在ziketi3环境启动；QGIS与MCP容器配置应使用示例路径。",
        "最终ZIP和视频分别不超过200MB，并在截止日前完成培养单位资格审核。",
    ):
        add_bullet(doc, item)

    add_heading(doc, "5 当前已知风险与应对", 1)
    add_table(
        doc,
        ["风险", "影响", "应对措施"],
        [
            ("32B模型推理耗时波动", "智能解释等待时间增加", "模型阶段超时回退；继续进行量化、缓存和上下文裁剪"),
            ("规则仍需专家复核", "阈值或适用场景可能偏差", "保持规则来源可追溯，建立专家审核与版本发布流程"),
            ("样例项目数量有限", "泛化效果尚不能充分证明", "扩充不同地区、不同图层结构的脱敏数据集"),
            ("标准版权和工程数据安全", "不当分发可能产生合规风险", "ZIP仅保留结构化摘要、来源清单和授权样例；敏感数据不提交"),
        ],
        [2400, 2500, 4160],
        9.3,
    )
    doc.core_properties.title = "初赛提交材料清单与技术可行性说明"
    doc.core_properties.author = TEAM_NAME
    doc.save(output_path)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    intro_path = OUTPUT_DIR / f"{TEAM_NAME}_{PROJECT_NAME}_参赛作品简介.docx"
    project_path = OUTPUT_DIR / f"{TEAM_NAME}_{PROJECT_NAME}_项目文档.docx"
    auxiliary_path = OUTPUT_DIR / f"{TEAM_NAME}_{PROJECT_NAME}_其他辅助材料说明.docx"
    create_intro_document(intro_path)
    create_project_document(project_path)
    create_auxiliary_document(auxiliary_path)
    print(f"intro_chars={len(INTRO_TEXT)}")
    for path in (intro_path, project_path, auxiliary_path):
        print(path)


if __name__ == "__main__":
    main()
